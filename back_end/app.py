import csv
import io
import os
import re
import secrets
import unicodedata
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import mysql.connector
from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, request, send_file, send_from_directory, session, url_for
from mysql.connector import IntegrityError
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


ROOT_DIR = Path(__file__).resolve().parent.parent
FRONT_END_DIR = ROOT_DIR / "front-end"
load_dotenv(ROOT_DIR / "back_end" / ".env")

TEXT_FIELDS = (
    "code_tresor_pay",
    "email",
    "nom",
    "prenoms",
    "sexe",
    "lieu_naissance",
    "nationalite",
    "telephone",
    "nature_piece",
    "numero_piece",
    "commune",
    "ville",
    "adresse",
    "serie_bac",
    "numero_bac",
    "numero_table",
    "mention",
    "choix_1_filiere",
    "choix_2_filiere",
    "tuteur1_nom",
    "tuteur1_contact",
    "tuteur1_lien",
    "tuteur1_residence",
    "tuteur2_nom",
    "tuteur2_contact",
    "tuteur2_lien",
    "tuteur2_residence",
)
DECIMAL_FIELDS = (
    "moyenne_bac",
    "note_math_bac",
    "note_physique_bac",
    "note_francais_bac",
    "note_anglais_bac",
)
DOCUMENT_TYPES = {
    "attestation": {"pdf", "jpg", "jpeg", "png"},
    "releve_bac": {"pdf", "jpg", "jpeg", "png"},
    "cni": {"pdf", "jpg", "jpeg", "png"},
    "bulletins_seconde": {"pdf", "jpg", "jpeg", "png"},
    "bulletins_premiere": {"pdf", "jpg", "jpeg", "png"},
    "bulletins_terminale": {"pdf", "jpg", "jpeg", "png"},
    "photo": {"jpg", "jpeg", "png"},
    "lettre_motivation": {"pdf", "doc", "docx"},
    "acte_naissance": {"pdf", "jpg", "jpeg", "png"},
    "cv": {"pdf", "doc", "docx"},
}
REQUIRED_CANDIDATE_FIELDS = (
    "code_tresor_pay",
    "email",
    "nom",
    "prenoms",
    "sexe",
    "date_naissance",
    "telephone",
    "annee_bac",
    "serie_bac",
    "choix_1_filiere",
)


def get_connection():
    return mysql.connector.connect(
        host=os.getenv("DB_HOST", "sql312.infinityfree.com"),
        port=int(os.getenv("DB_PORT", "3306")),
        user=os.getenv("DB_USER", "if0_42066052"),
        password=os.getenv("DB_PASSWORD", "LeKageNo122006"),
        database=os.getenv("DB_NAME", "if0_42066052_emsp2"),
        charset="utf8mb4",
    )


def create_app(test_config=None):
    app = Flask(__name__)
    upload_folder = Path(os.getenv("UPLOAD_FOLDER", "uploads"))
    if not upload_folder.is_absolute():
        upload_folder = ROOT_DIR / "back_end" / upload_folder
    app.config.from_mapping(
        SECRET_KEY=os.getenv("FLASK_SECRET_KEY") or secrets.token_hex(32),
        UPLOAD_FOLDER=str(upload_folder),
        MAX_CONTENT_LENGTH=64 * 1024 * 1024,
        MAX_DOCUMENT_SIZE=5 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "false").lower()
        == "true",
    )
    if test_config:
        app.config.update(test_config)

    @app.get("/")
    def index():
        return send_from_directory(FRONT_END_DIR, "accueil.html")

    @app.get("/connexion")
    def connexion_page():
        return send_from_directory(FRONT_END_DIR, "connexion.html")

    @app.get("/candidature")
    def candidature_page():
        # Toute tentative directe depuis un CTA public doit passer par la connexion.
        # Le lien "Créer un compte" de la page de connexion utilise le paramètre
        # `inscription=1` pour permettre à un nouveau candidat d'ouvrir le formulaire.
        inscription_request = request.args.get("inscription") == "1"
        if not session.get("numero_dossier") and not inscription_request:
            return redirect(url_for("connexion_page"))
        if inscription_request and not session.get("numero_dossier"):
            connection = None
            cursor = None
            try:
                connection = get_connection()
                cursor = connection.cursor(dictionary=True)
                if not registrations_are_open(cursor):
                    return redirect(url_for("connexion_page", inscriptions="ferme"))
            except mysql.connector.Error:
                app.logger.exception("Impossible de vérifier l'état des inscriptions")
            finally:
                if cursor:
                    cursor.close()
                if connection:
                    connection.close()
        return send_from_directory(FRONT_END_DIR, "candidature.html")

    @app.get("/contacts")
    def contacts_page():
        return send_from_directory(FRONT_END_DIR, "contacts.html")

    @app.get("/dashboard")
    def dashboard_page():
        if session.get("role") != "admin":
            return redirect(url_for("connexion_page"))
        return send_from_directory(FRONT_END_DIR, "dashboard.html")

    @app.get("/css/<path:filename>")
    def frontend_css(filename):
        return send_from_directory(FRONT_END_DIR / "css", filename)

    @app.get("/js/<path:filename>")
    def frontend_js(filename):
        return send_from_directory(FRONT_END_DIR / "js", filename)

    @app.get("/assets/<path:filename>")
    def frontend_assets(filename):
        return send_from_directory(FRONT_END_DIR / "assets", filename)

    def is_admin_login(identifier, password):
        admin_email = os.getenv("ADMIN_EMAIL", "admin@emsp.ci").strip().lower()
        if app.config.get("TESTING") and identifier.lower() != admin_email:
            return False
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT password_hash FROM admin_user "
                "WHERE email = %s AND actif = 1 LIMIT 1",
                (identifier.lower(),),
            )
            admin = cursor.fetchone()
            return bool(
                admin and admin.get("password_hash")
                and check_password_hash(admin["password_hash"], password)
            )
        except (mysql.connector.Error, ValueError):
            # Le hash dans .env permet de conserver un secours local si la table
            # n'est pas encore initialisée, sans exposer le mot de passe en clair.
            if identifier.lower() != admin_email:
                return False
            admin_password_hash = os.getenv("ADMIN_PASSWORD_HASH", "").strip()
            if not admin_password_hash:
                return False
            try:
                return check_password_hash(admin_password_hash, password)
            except ValueError:
                return False
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    def admin_required():
        return session.get("role") == "admin"

    def super_admin_required():
        return admin_required() and session.get("admin_email") == os.getenv(
            "ADMIN_EMAIL", "admin@emsp.ci"
        ).strip().lower()

    def modifications_are_open(cursor):
        cursor.execute(
            "SELECT setting_value FROM platform_setting "
            "WHERE setting_key = 'candidate_edits_open' LIMIT 1"
        )
        setting = cursor.fetchone()
        if not setting:
            return True
        if isinstance(setting, dict):
            value = setting.get("setting_value")
            if value is None:
                return True
        else:
            value = setting[0]
        return str(value).lower() in {"1", "true", "yes", "open"}

    def registrations_are_open(cursor):
        return str(platform_setting(cursor, "candidate_registrations_open", "true")).lower() in {
            "1", "true", "yes", "open"
        }

    def platform_setting(cursor, key, default=""):
        cursor.execute(
            "SELECT setting_value FROM platform_setting WHERE setting_key = %s LIMIT 1",
            (key,),
        )
        setting = cursor.fetchone()
        if not setting:
            return default
        return setting.get("setting_value", default) if isinstance(setting, dict) else setting[0]

    @app.get("/api/health")
    def health():
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute("SELECT 1")
            cursor.fetchone()
            return jsonify(status="ok", database="connected")
        except mysql.connector.Error:
            return jsonify(status="error", database="unavailable"), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/candidatures")
    def create_candidature():
        if not app.config.get("TESTING"):
            gate_connection = get_connection()
            gate_cursor = gate_connection.cursor(dictionary=True)
            try:
                if not registrations_are_open(gate_cursor):
                    return jsonify(error="Les inscriptions sont actuellement clôturées par l’administration."), 423
            finally:
                gate_cursor.close()
                gate_connection.close()
        form = request.form
        missing = [
            field for field in REQUIRED_CANDIDATE_FIELDS
            if not form.get(field, "").strip()
        ]
        password = form.get("password", "")
        if missing or not password:
            if not password:
                missing.append("password")
            return jsonify(error="Champs obligatoires manquants.", fields=missing), 400
        if form.get("confirmation") != "on":
            return jsonify(error="Veuillez confirmer l'exactitude des informations."), 400

        email = form.get("email", "").strip().lower()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            return jsonify(error="Adresse email invalide.", fields=["email"]), 400
        if len(password) < 8:
            return jsonify(
                error="Le mot de passe doit contenir au moins 8 caractères.",
                fields=["password"],
            ), 400

        values = {field: form.get(field, "").strip() or None for field in TEXT_FIELDS}
        values["email"] = email
        values["MDP"] = generate_password_hash(password)

        try:
            values["date_naissance"] = date.fromisoformat(
                form["date_naissance"]
            ).isoformat()
            values["annee_bac"] = int(form["annee_bac"])
            if not 1901 <= values["annee_bac"] <= 2155:
                raise ValueError
            for field in DECIMAL_FIELDS:
                raw_value = form.get(field, "").strip()
                values[field] = Decimal(raw_value) if raw_value else None
                if values[field] is not None and not (
                    Decimal("0") <= values[field] <= Decimal("20")
                ):
                    raise ValueError
        except (ValueError, InvalidOperation):
            return jsonify(error="Date, année du BAC ou note invalide."), 400

        uploaded_files = []
        for document_type, upload in request.files.items():
            if document_type not in DOCUMENT_TYPES:
                return jsonify(error="Type de document non autorisé."), 400
            if not upload or not upload.filename:
                continue
            safe_name = secure_filename(upload.filename)
            extension = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
            if not safe_name or extension not in DOCUMENT_TYPES[document_type]:
                return jsonify(
                    error=f"Format de fichier invalide : {document_type}."
                ), 400
            upload.stream.seek(0, os.SEEK_END)
            file_size = upload.stream.tell()
            upload.stream.seek(0)
            if file_size > app.config["MAX_DOCUMENT_SIZE"]:
                return jsonify(error=f"Fichier trop volumineux : {document_type}."), 400
            stored_name = f"{document_type}_{secrets.token_hex(16)}.{extension}"
            uploaded_files.append((document_type, upload, safe_name, stored_name))

        if not app.config.get("TESTING"):
            check_connection = None
            check_cursor = None
            try:
                check_connection = get_connection()
                check_cursor = check_connection.cursor(dictionary=True)
                check_cursor.execute(
                    "SELECT statut_candidature FROM candidature "
                    "WHERE email = %s OR (LOWER(nom) = LOWER(%s) "
                    "AND LOWER(prenoms) = LOWER(%s))",
                    (email, values["nom"], values["prenoms"]),
                )
                existing_rows = check_cursor.fetchall()
                if any(row.get("statut_candidature") == "rejete" for row in existing_rows):
                    return jsonify(
                        error="Cette candidature a déjà été refusée. Vous ne pouvez plus créer de dossier."
                    ), 403
            finally:
                if check_cursor:
                    check_cursor.close()
                if check_connection:
                    check_connection.close()

        dossier = f"CDT_{secrets.token_hex(6).upper()}"
        dossier_dir = Path(app.config["UPLOAD_FOLDER"]) / dossier
        saved_paths = []
        connection = None
        cursor = None
        try:
            dossier_dir.mkdir(parents=True, exist_ok=True)
            for _, upload, _, stored_name in uploaded_files:
                path = dossier_dir / stored_name
                upload.save(path)
                saved_paths.append(path)

            connection = get_connection()
            cursor = connection.cursor()
            columns = [
                "numero_dossier", *TEXT_FIELDS, "MDP", "date_naissance",
                "annee_bac", *DECIMAL_FIELDS,
            ]
            insert_values = [
                dossier, *(values[field] for field in TEXT_FIELDS), values["MDP"],
                values["date_naissance"], values["annee_bac"],
                *(values[field] for field in DECIMAL_FIELDS),
            ]
            cursor.execute(
                f"INSERT INTO candidature ({', '.join(columns)}) "
                f"VALUES ({', '.join(['%s'] * len(columns))})",
                insert_values,
            )
            for document_type, _, original_name, stored_name in uploaded_files:
                relative_path = f"{dossier}/{stored_name}"
                cursor.execute(
                    "INSERT INTO candidature_document "
                    "(numero_dossier, type_document, nom_fichier, chemin_fichier) "
                    "VALUES (%s, %s, %s, %s)",
                    (dossier, document_type, original_name, relative_path),
                )
            connection.commit()
        except IntegrityError:
            if connection:
                connection.rollback()
            for path in saved_paths:
                path.unlink(missing_ok=True)
            if dossier_dir.exists() and not any(dossier_dir.iterdir()):
                dossier_dir.rmdir()
            return jsonify(error="Email ou code Trésor Pay déjà utilisé."), 409
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            for path in saved_paths:
                path.unlink(missing_ok=True)
            if dossier_dir.exists() and not any(dossier_dir.iterdir()):
                dossier_dir.rmdir()
            app.logger.exception("Erreur lors de l'enregistrement de la candidature")
            return jsonify(error="La base de données est indisponible."), 503
        except OSError:
            for path in saved_paths:
                path.unlink(missing_ok=True)
            if dossier_dir.exists() and not any(dossier_dir.iterdir()):
                dossier_dir.rmdir()
            app.logger.exception("Erreur lors du stockage des documents")
            return jsonify(error="Impossible d'enregistrer les documents."), 500
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

        return jsonify(message="Candidature enregistrée.", numero_dossier=dossier), 201

    @app.post("/api/auth/login")
    def login():
        data = request.get_json(silent=True) or {}
        identifier = str(data.get("identifiant", "")).strip()
        password = str(data.get("password", ""))
        if not identifier or not password:
            return jsonify(error="Identifiant et mot de passe requis."), 400

        if is_admin_login(identifier, password):
            session.clear()
            session["role"] = "admin"
            session["admin_email"] = identifier.lower()
            session["is_super_admin"] = identifier.lower() == os.getenv(
                "ADMIN_EMAIL", "admin@emsp.ci"
            ).strip().lower()
            return jsonify(role="admin", redirect="/dashboard")

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT numero_dossier, email, nom, prenoms, MDP, statut_candidature FROM candidature "
                "WHERE email = %s OR numero_dossier = %s LIMIT 1",
                (identifier.lower(), identifier),
            )
            candidate = cursor.fetchone()
        except mysql.connector.Error:
            app.logger.exception("Erreur lors de la connexion candidat")
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

        if candidate and candidate.get("statut_candidature") == "rejete":
            return jsonify(error="Votre candidature a été refusée. Accès interdit."), 403
        if not candidate or not check_password_hash(candidate["MDP"], password):
            return jsonify(error="Identifiant ou mot de passe incorrect."), 401
        session.clear()
        session["role"] = "candidate"
        session["numero_dossier"] = candidate["numero_dossier"]
        return jsonify(
            role="candidate",
            numero_dossier=candidate["numero_dossier"],
            nom=candidate["nom"],
            prenoms=candidate["prenoms"],
        )

    @app.get("/api/auth/me")
    def current_candidate():
        dossier = session.get("numero_dossier")
        if not dossier:
            return jsonify(error="Authentification requise."), 401
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT * FROM candidature WHERE numero_dossier = %s",
                (dossier,),
            )
            candidate = cursor.fetchone()
            if candidate:
                candidate.pop("MDP", None)
                candidate["modifications_ouvertes"] = modifications_are_open(cursor)
                candidate["statut_candidature"] = candidate.get("statut_candidature") or (
                    "valide" if candidate.get("dossier_valide") else "en_traitement"
                )
                cursor.execute(
                    "SELECT 1 FROM candidature_document "
                    "WHERE numero_dossier = %s AND type_document = 'convocation' "
                    "LIMIT 1",
                    (dossier,),
                )
                candidate["convocation_disponible"] = bool(cursor.fetchone())
                cursor.execute(
                    "SELECT type_document FROM candidature_document WHERE numero_dossier = %s",
                    (dossier,),
                )
                document_rows = cursor.fetchall() if hasattr(cursor, "fetchall") else []
                candidate["documents_deposes"] = [row["type_document"] for row in document_rows]
        except mysql.connector.Error:
            app.logger.exception("Erreur lors de la lecture du profil candidat")
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()
        if not candidate:
            session.clear()
            return jsonify(error="Candidat introuvable."), 404
        return jsonify(candidate)

    @app.put("/api/candidatures/me")
    def update_current_candidate():
        if session.get("role") != "candidate" or not session.get("numero_dossier"):
            return jsonify(error="Authentification candidat requise."), 401

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            if not modifications_are_open(cursor):
                return jsonify(error="Les modifications de dossiers sont clôturées."), 423
            cursor.execute(
                "SELECT dossier_valide, statut_candidature FROM candidature WHERE numero_dossier = %s LIMIT 1",
                (session["numero_dossier"],),
            )
            current = cursor.fetchone()
            if not current:
                return jsonify(error="Candidat introuvable."), 404
            if current["dossier_valide"] or current.get("statut_candidature") == "rejete":
                return jsonify(error="Votre dossier est validé et ne peut plus être modifié."), 423

            form = request.form
            updates = {}
            for field in TEXT_FIELDS:
                if field in form and field != "email":
                    updates[field] = form.get(field, "").strip() or None
            if "email" in form:
                email = form.get("email", "").strip().lower()
                if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
                    return jsonify(error="Adresse email invalide.", fields=["email"]), 400
                updates["email"] = email
            if "date_naissance" in form:
                updates["date_naissance"] = date.fromisoformat(form["date_naissance"]).isoformat()
            if "annee_bac" in form:
                updates["annee_bac"] = int(form["annee_bac"])
            for field in DECIMAL_FIELDS:
                if field in form:
                    raw_value = form.get(field, "").strip()
                    updates[field] = Decimal(raw_value) if raw_value else None
                    if updates[field] is not None and not Decimal("0") <= updates[field] <= Decimal("20"):
                        raise ValueError
            password = form.get("password", "")
            if password:
                if len(password) < 8:
                    return jsonify(error="Le mot de passe doit contenir au moins 8 caractères."), 400
                updates["MDP"] = generate_password_hash(password)
            if not updates:
                return jsonify(error="Aucune donnée à modifier."), 400

            assignments = ", ".join(f"{field} = %s" for field in updates)
            cursor.execute(
                f"UPDATE candidature SET {assignments} WHERE numero_dossier = %s",
                [*updates.values(), session["numero_dossier"]],
            )
            connection.commit()
            return jsonify(message="Vos informations ont été mises à jour.")
        except (ValueError, InvalidOperation):
            return jsonify(error="Date, année du BAC ou note invalide."), 400
        except IntegrityError:
            if connection:
                connection.rollback()
            return jsonify(error="Email ou code Trésor Pay déjà utilisé."), 409
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/auth/logout")
    def logout():
        session.clear()
        return jsonify(message="Session terminée.")

    @app.post("/api/auth/reset-password")
    def reset_password():
        data = request.get_json(silent=True) or {}
        identifier = str(data.get("identifiant", "")).strip()
        new_password = str(data.get("password", ""))
        confirmation = str(data.get("confirmation", ""))
        personal_answer = str(data.get("reponse_personnelle", "")).strip()
        if not identifier or not new_password or new_password != confirmation:
            return jsonify(error="Identifiant et confirmation du mot de passe requis."), 400
        if len(new_password) < 8:
            return jsonify(error="Le mot de passe doit contenir au moins 8 caractères."), 400

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            password_hash = generate_password_hash(new_password)
            cursor.execute("SELECT id FROM admin_user WHERE email = %s AND actif = 1 LIMIT 1", (identifier.lower(),))
            is_admin_account = cursor.fetchone()
            if is_admin_account and not super_admin_required():
                return jsonify(error="Seul le super administrateur peut modifier un mot de passe administrateur."), 403
            cursor.execute(
                "UPDATE admin_user SET password_hash = %s "
                "WHERE email = %s AND actif = 1",
                (password_hash, identifier.lower()),
            )
            updated = cursor.rowcount > 0
            if not updated:
                cursor.execute(
                    "SELECT lieu_naissance, statut_candidature FROM candidature "
                    "WHERE email = %s OR numero_dossier = %s",
                    (identifier.lower(), identifier),
                )
                candidate = cursor.fetchone()
                if not candidate:
                    return jsonify(error="Email ou numéro de dossier introuvable."), 404
                if candidate.get("statut_candidature") == "rejete":
                    return jsonify(error="Cette candidature est refusée et l’accès est bloqué."), 403
                def normalize_answer(value):
                    return " ".join(unicodedata.normalize("NFKD", str(value or ""))
                                   .encode("ascii", "ignore").decode().lower().split())
                if not personal_answer or normalize_answer(personal_answer) != normalize_answer(candidate.get("lieu_naissance")):
                    connection.rollback()
                    return jsonify(error="La réponse à la question personnelle est incorrecte."), 403
                cursor.execute(
                    "UPDATE candidature SET MDP = %s WHERE email = %s OR numero_dossier = %s",
                    (password_hash, identifier.lower(), identifier),
                )
                updated = cursor.rowcount > 0
            if not updated:
                connection.rollback()
                return jsonify(error="Email ou numéro de dossier introuvable."), 404
            connection.commit()
            session.clear()
            return jsonify(message="Mot de passe modifié avec succès.")
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.get("/api/candidatures/me/convocation")
    def candidate_convocation():
        if session.get("role") != "candidate" or not session.get("numero_dossier"):
            return jsonify(error="Authentification candidat requise."), 401
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT chemin_fichier, nom_fichier FROM candidature_document "
                "WHERE numero_dossier = %s AND type_document = 'convocation' LIMIT 1",
                (session["numero_dossier"],),
            )
            document = cursor.fetchone()
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()
        if not document:
            return jsonify(error="Aucune convocation n’est disponible."), 404
        upload_root = Path(app.config["UPLOAD_FOLDER"]).resolve()
        document_path = (upload_root / document["chemin_fichier"]).resolve()
        if upload_root not in document_path.parents or not document_path.is_file():
            return jsonify(error="Convocation indisponible."), 404
        return send_from_directory(
            document_path.parent,
            document_path.name,
            as_attachment=True,
            download_name=secure_filename(document["nom_fichier"]),
        )

    @app.get("/api/admin/dashboard")
    def admin_dashboard_data():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT COUNT(*) AS total, "
                "SUM(CASE WHEN COALESCE(statut_candidature, IF(dossier_valide = 1, 'valide', 'en_traitement')) = 'valide' THEN 1 ELSE 0 END) AS valides, "
                "SUM(CASE WHEN statut_candidature = 'rejete' THEN 1 ELSE 0 END) AS rejetes, "
                "SUM(CASE WHEN admis_concours = 1 THEN 1 ELSE 0 END) AS admis "
                "FROM candidature"
            )
            counts = cursor.fetchone() or {}
            search = request.args.get("q", "").strip()
            query = (
                "SELECT numero_dossier, nom, prenoms, email, choix_1_filiere, "
                "dossier_valide, statut_candidature, admis_concours "
                "FROM candidature"
            )
            params = []
            status_filter = request.args.get("status", "").strip().lower()
            if search:
                query += " WHERE numero_dossier LIKE %s OR email LIKE %s OR nom LIKE %s OR prenoms LIKE %s"
                pattern = f"%{search}%"
                params = [pattern, pattern, pattern, pattern]
            if status_filter in {"en_traitement", "valide", "rejete"}:
                query += " AND " if " WHERE " in query else " WHERE "
                query += " COALESCE(statut_candidature, IF(dossier_valide = 1, 'valide', 'en_traitement')) = %s"
                params.append(status_filter)
            query += " ORDER BY numero_dossier DESC LIMIT 100"
            cursor.execute(query, params)
            recent = cursor.fetchall() or []
            modifications_open = modifications_are_open(cursor)
            registrations_open = registrations_are_open(cursor)
            composition_date = platform_setting(cursor, "composition_date")
            composition_centre = platform_setting(cursor, "composition_centre")
            return jsonify(
                counts={
                    "total": int(counts.get("total") or 0),
                    "valides": int(counts.get("valides") or 0),
                    "admis": int(counts.get("admis") or 0),
                    "rejetes": int(counts.get("rejetes") or 0),
                },
                recent=recent,
                modifications_ouvertes=modifications_open,
                inscriptions_ouvertes=registrations_open,
                composition={"date": composition_date, "centre": composition_centre},
            )
        except mysql.connector.Error:
            app.logger.exception("Erreur lors du chargement du dashboard admin")
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/convocations/generate")
    def admin_generate_convocations():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        connection = None
        try:
            from generate_convocations import generate_all_convocations

            connection = get_connection()
            result = generate_all_convocations(connection, app.config["UPLOAD_FOLDER"])
            return jsonify(
                message="Convocations générées pour les dossiers validés.",
                **result,
            )
        except FileNotFoundError as error:
            return jsonify(error=str(error)), 500
        except RuntimeError as error:
            app.logger.exception("Erreur lors de la conversion des convocations")
            return jsonify(error=str(error)), 500
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            app.logger.exception("Erreur lors de la génération des convocations")
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if connection:
                connection.close()

    @app.get("/api/admin/candidates/export.csv")
    def admin_export_candidates():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT numero_dossier, nom, prenoms, email, telephone, "
                "choix_1_filiere, choix_2_filiere, "
                "COALESCE(statut_candidature, IF(dossier_valide = 1, 'valide', 'en_traitement')) AS statut, "
                "date_naissance FROM candidature ORDER BY numero_dossier"
            )
            output = io.StringIO(newline="")
            # Excel configuré en français attend généralement le point-virgule.
            writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
            writer.writerow([
                "Numéro de dossier", "Nom", "Prénoms", "Email", "Téléphone",
                "Choix 1", "Choix 2", "Statut", "Date de naissance",
            ])
            for candidate in cursor.fetchall() or []:
                writer.writerow([
                    candidate.get("numero_dossier"), candidate.get("nom"),
                    candidate.get("prenoms"), candidate.get("email"),
                    candidate.get("telephone"), candidate.get("choix_1_filiere"),
                    candidate.get("choix_2_filiere"), candidate.get("statut"),
                    candidate.get("date_naissance"),
                ])
            response = app.response_class(
                output.getvalue().encode("utf-8-sig"),
                mimetype="text/csv; charset=utf-8",
            )
            response.headers["Content-Disposition"] = "attachment; filename=candidats_emsp.csv"
            return response
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.get("/api/admin/candidates/export.xlsx")
    def admin_export_candidates_excel():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        connection = None
        cursor = None
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter

            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT numero_dossier, nom, prenoms, email, telephone, "
                "choix_1_filiere, choix_2_filiere, "
                "COALESCE(statut_candidature, IF(dossier_valide = 1, 'valide', 'en_traitement')) AS statut, "
                "date_naissance FROM candidature ORDER BY numero_dossier"
            )
            headers = [
                "Numéro de dossier", "Nom", "Prénoms", "Email", "Téléphone",
                "Choix 1", "Choix 2", "Statut", "Date de naissance",
            ]
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Candidats"
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="155B3A")
            for candidate in cursor.fetchall() or []:
                sheet.append([
                    candidate.get("numero_dossier"), candidate.get("nom"),
                    candidate.get("prenoms"), candidate.get("email"),
                    candidate.get("telephone"), candidate.get("choix_1_filiere"),
                    candidate.get("choix_2_filiere"), candidate.get("statut"),
                    candidate.get("date_naissance"),
                ])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                sheet.column_dimensions[get_column_letter(column[0].column)].width = width
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)
            return send_file(
                output, as_attachment=True, download_name="candidats_emsp.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except ImportError:
            return jsonify(error="Le module openpyxl est requis pour l'export Excel."), 500
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.get("/api/admin/candidates/<string:dossier>")
    def admin_candidate_details(dossier):
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT * FROM candidature WHERE numero_dossier = %s LIMIT 1",
                (dossier,),
            )
            candidate = cursor.fetchone()
            if candidate:
                candidate["modifications_ouvertes"] = modifications_are_open(cursor)
            if not candidate:
                return jsonify(error="Candidat introuvable."), 404

            candidate.pop("MDP", None)
            for key, value in list(candidate.items()):
                if hasattr(value, "isoformat"):
                    candidate[key] = value.isoformat()
                elif isinstance(value, Decimal):
                    candidate[key] = str(value)

            cursor.execute(
                "SELECT id, type_document, nom_fichier, date_ajout "
                "FROM candidature_document WHERE numero_dossier = %s "
                "ORDER BY date_ajout DESC",
                (dossier,),
            )
            documents = cursor.fetchall() or []
            for document in documents:
                if hasattr(document.get("date_ajout"), "isoformat"):
                    document["date_ajout"] = document["date_ajout"].isoformat()
                document["download_url"] = url_for(
                    "admin_download_document", dossier=dossier, document_id=document["id"]
                )
            return jsonify(candidate=candidate, documents=documents)
        except mysql.connector.Error:
            app.logger.exception("Erreur lors de la lecture du dossier candidat")
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/candidates/<string:dossier>/results")
    def admin_update_candidate_results(dossier):
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        note_fields = (
            "note_francais_compo", "note_math_compo",
            "note_anglais_compo", "note_psycho_compo",
        )
        notes = {}
        try:
            for field in note_fields:
                raw = data.get(field)
                notes[field] = None if raw in (None, "") else Decimal(str(raw))
                if notes[field] is not None and not Decimal("0") <= notes[field] <= Decimal("20"):
                    raise ValueError
        except (InvalidOperation, ValueError):
            return jsonify(error="Chaque note doit être comprise entre 0 et 20."), 400
        admitted = bool(data.get("admis_concours"))
        filiere = str(data.get("filiere_formation", "")).strip() or None
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            assignments = ", ".join(f"{field} = %s" for field in note_fields)
            cursor.execute(
                f"UPDATE candidature SET {assignments}, admis_concours = %s, "
                "filiere_formation = %s WHERE numero_dossier = %s",
                [*(notes[field] for field in note_fields), admitted, filiere, dossier],
            )
            if cursor.rowcount == 0:
                return jsonify(error="Candidat introuvable."), 404
            connection.commit()
            return jsonify(message="Résultats enregistrés.", admis_concours=admitted)
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.get("/api/admin/results/template.xlsx")
    def admin_results_template():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        from results_excel import template_bytes

        return send_file(
            template_bytes(), as_attachment=True, download_name="modele_resultats_emsp.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.post("/api/admin/results/import")
    def admin_import_results():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        upload = request.files.get("results_file")
        if not upload or not upload.filename.lower().endswith((".xlsx", ".xlsm")):
            return jsonify(error="Veuillez sélectionner un fichier Excel .xlsx."), 400
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook["Resultats"] if "Resultats" in workbook.sheetnames else workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return jsonify(error="Le fichier Excel est vide."), 400
            headers = [str(value or "").strip() for value in rows[0]]
            required = {"numero_dossier", "email", "publier"}
            if not required.issubset(headers):
                return jsonify(error="Colonnes obligatoires manquantes dans le fichier Excel."), 400
            positions = {header: index for index, header in enumerate(headers)}
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            updated = []
            errors = []
            for line_number, row in enumerate(rows[1:], start=2):
                values = list(row) + [None] * (len(headers) - len(row))
                publish = str(values[positions["publier"]] or "").strip().lower() in {"oui", "o", "yes", "1", "true"}
                if not publish:
                    continue
                dossier = str(values[positions["numero_dossier"]] or "").strip()
                email = str(values[positions["email"]] or "").strip().lower()
                if not dossier and not email:
                    errors.append(f"Ligne {line_number}: dossier ou email obligatoire.")
                    continue
                if dossier:
                    cursor.execute("SELECT numero_dossier FROM candidature WHERE numero_dossier = %s LIMIT 1", (dossier,))
                else:
                    cursor.execute("SELECT numero_dossier FROM candidature WHERE email = %s LIMIT 1", (email,))
                candidate = cursor.fetchone()
                if not candidate:
                    errors.append(f"Ligne {line_number}: candidat introuvable.")
                    continue
                target = candidate["numero_dossier"]
                note_fields = ("note_francais_compo", "note_math_compo", "note_anglais_compo", "note_psycho_compo")
                note_values = []
                invalid = False
                for field in note_fields:
                    index = positions.get(field)
                    value = values[index] if index is not None else None
                    if value in (None, ""):
                        note_values.append(None)
                    else:
                        try:
                            parsed = Decimal(str(value))
                            if not Decimal("0") <= parsed <= Decimal("20"):
                                raise ValueError
                            note_values.append(parsed)
                        except (InvalidOperation, ValueError):
                            invalid = True
                if invalid:
                    errors.append(f"Ligne {line_number}: note invalide.")
                    continue
                admitted_index = positions.get("admis_concours")
                admitted = str(values[admitted_index] if admitted_index is not None else "").lower() in {"oui", "o", "yes", "1", "true"}
                filiere_index = positions.get("filiere_formation")
                filiere = values[filiere_index] if filiere_index is not None else None
                cursor.execute(
                    "UPDATE candidature SET note_francais_compo=%s, note_math_compo=%s, "
                    "note_anglais_compo=%s, note_psycho_compo=%s, admis_concours=%s, filiere_formation=%s "
                    "WHERE numero_dossier=%s",
                    (*note_values, admitted, str(filiere).strip() if filiere else None, target),
                )
                updated.append(target)
            connection.commit()
            cursor.close()
            connection.close()
            return jsonify(message="Résultats importés et publiés.", updated=updated, errors=errors)
        except ImportError:
            return jsonify(error="Le module openpyxl est requis pour l'import Excel."), 500
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503

    @app.get("/api/admin/candidates/<string:dossier>/documents/<int:document_id>")
    def admin_download_document(dossier, document_id):
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute(
                "SELECT chemin_fichier, nom_fichier FROM candidature_document "
                "WHERE id = %s AND numero_dossier = %s LIMIT 1",
                (document_id, dossier),
            )
            document = cursor.fetchone()
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

        if not document:
            return jsonify(error="Document introuvable."), 404
        upload_root = Path(app.config["UPLOAD_FOLDER"]).resolve()
        document_path = (upload_root / document["chemin_fichier"]).resolve()
        if upload_root not in document_path.parents or not document_path.is_file():
            return jsonify(error="Document indisponible."), 404
        return send_from_directory(
            document_path.parent,
            document_path.name,
            as_attachment=True,
            download_name=secure_filename(document["nom_fichier"]),
        )

    @app.post("/api/admin/candidates/<string:dossier>/validate")
    def admin_validate_candidate(dossier):
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        status = str(data.get("status", "valide")).strip().lower()
        if status not in {"en_traitement", "valide", "rejete"}:
            return jsonify(error="Statut de candidature invalide."), 400
        validated = status == "valide"
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "UPDATE candidature SET dossier_valide = %s, statut_candidature = %s "
                "WHERE numero_dossier = %s",
                (validated, status, dossier),
            )
            if cursor.rowcount == 0:
                return jsonify(error="Candidat introuvable."), 404
            connection.commit()
            return jsonify(message="Statut du dossier mis à jour.", dossier_valide=validated, statut=status)
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/settings/modifications")
    def admin_update_modifications_setting():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        is_open = bool(data.get("open", False))
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "INSERT INTO platform_setting (setting_key, setting_value) "
                "VALUES ('candidate_edits_open', %s) "
                "ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value)",
                ("true" if is_open else "false",),
            )
            connection.commit()
            return jsonify(message="Paramètre des modifications mis à jour.", ouvert=is_open)
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/settings/registrations")
    def admin_update_registrations_setting():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        is_open = bool(data.get("open", False))
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "INSERT INTO platform_setting (setting_key, setting_value) VALUES "
                "('candidate_registrations_open', %s) "
                "ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value)",
                ("true" if is_open else "false",),
            )
            connection.commit()
            return jsonify(message="Paramètre des inscriptions mis à jour.", ouvert=is_open)
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/settings/composition")
    def admin_update_composition_setting():
        if not admin_required():
            return jsonify(error="Accès administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        composition_date = str(data.get("date", "")).strip()
        composition_centre = str(data.get("centre", "")).strip()
        try:
            date.fromisoformat(composition_date)
        except ValueError:
            return jsonify(error="La date de composition est invalide."), 400
        if not composition_centre:
            return jsonify(error="Le centre de composition est obligatoire."), 400
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            for key, value in (("composition_date", composition_date), ("composition_centre", composition_centre)):
                cursor.execute(
                    "INSERT INTO platform_setting (setting_key, setting_value) VALUES (%s, %s) "
                    "ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value)",
                    (key, value),
                )
            connection.commit()
            return jsonify(message="Paramètres de composition enregistrés.", date=composition_date, centre=composition_centre)
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/admins")
    def admin_create_admin():
        if not super_admin_required():
            return jsonify(error="Seul le super administrateur peut gérer les comptes administrateurs."), 403
        data = request.get_json(silent=True) or {}
        email = str(data.get("email", "")).strip().lower()
        password = str(data.get("password", ""))
        name = str(data.get("nom", "Administrateur EMSP")).strip()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            return jsonify(error="Adresse email invalide."), 400
        if len(password) < 8:
            return jsonify(error="Le mot de passe doit contenir au moins 8 caractères."), 400
        if not name:
            return jsonify(error="Le nom de l’administrateur est requis."), 400

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "INSERT INTO admin_user (email, password_hash, nom, actif) "
                "VALUES (%s, %s, %s, TRUE)",
                (email, generate_password_hash(password), name),
            )
            connection.commit()
        except IntegrityError:
            if connection:
                connection.rollback()
            return jsonify(error="Cet email administrateur existe déjà."), 409
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()
        return jsonify(message="Administrateur créé."), 201

    @app.get("/api/admin/admins")
    def admin_list_admins():
        if not super_admin_required():
            return jsonify(error="Accès super administrateur requis."), 403
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor(dictionary=True)
            cursor.execute("SELECT id, email, nom, actif, date_creation FROM admin_user ORDER BY date_creation")
            admins = cursor.fetchall() or []
            for admin in admins:
                if hasattr(admin.get("date_creation"), "isoformat"):
                    admin["date_creation"] = admin["date_creation"].isoformat()
            return jsonify(admins=admins)
        except mysql.connector.Error:
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/admin/admins/<int:admin_id>/password")
    def admin_reset_admin_password(admin_id):
        if not super_admin_required():
            return jsonify(error="Accès super administrateur requis."), 403
        data = request.get_json(silent=True) or {}
        password = str(data.get("password", ""))
        if len(password) < 8:
            return jsonify(error="Le mot de passe doit contenir au moins 8 caractères."), 400
        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "UPDATE admin_user SET password_hash = %s WHERE id = %s AND actif = 1",
                (generate_password_hash(password), admin_id),
            )
            if cursor.rowcount == 0:
                return jsonify(error="Administrateur introuvable."), 404
            connection.commit()
            return jsonify(message="Mot de passe administrateur remplacé.")
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            return jsonify(error="La base de données est indisponible."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()

    @app.post("/api/contacts")
    def create_contact_message():
        data = request.get_json(silent=True) or request.form
        fields = {
            key: str(data.get(key, "")).strip()
            for key in ("nom", "email", "objet", "message")
        }
        if not all(fields.values()):
            return jsonify(error="Veuillez remplir tous les champs obligatoires."), 400
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", fields["email"]):
            return jsonify(error="Adresse email invalide."), 400
        telephone = str(data.get("telephone", "")).strip() or None

        connection = None
        cursor = None
        try:
            connection = get_connection()
            cursor = connection.cursor()
            cursor.execute(
                "INSERT INTO message_contact (nom, email, telephone, objet, message) "
                "VALUES (%s, %s, %s, %s, %s)",
                (
                    fields["nom"], fields["email"].lower(), telephone,
                    fields["objet"], fields["message"],
                ),
            )
            connection.commit()
        except mysql.connector.Error:
            if connection:
                connection.rollback()
            app.logger.exception("Erreur lors de l'enregistrement du message")
            return jsonify(error="Impossible d'enregistrer le message pour le moment."), 503
        finally:
            if cursor:
                cursor.close()
            if connection:
                connection.close()
        return jsonify(message="Votre message a été envoyé."), 201

    @app.post("/api/chat")
    def chat():
        data = request.get_json(silent=True) or {}
        message = str(data.get("message", "")).strip()
        if not message:
            return jsonify(error="Écrivez une question pour commencer."), 400

        normalized = "".join(
            character
            for character in unicodedata.normalize("NFD", message.lower())
            if unicodedata.category(character) != "Mn"
        )

        def contains(*terms):
            return any(term in normalized for term in terms)

        wants_links = contains("lien", "liens", "url", "ouvrir", "acceder", "page")
        links = []

        def add_links(*items):
            # Limite volontaire pour garder une réponse claire et exploitable.
            links.extend(items)
            del links[3:]

        if contains("bonjour", "salut", "bonsoir", "hello"):
            answer = (
                "Bonjour ! Je suis l’assistant EMSP. Je peux vous orienter dans "
                "la plateforme, les filières, la candidature et les documents."
            )
        elif contains("filiere", "formation", "etude", "dser", "fdig", "mdig", "lnum", "gare"):
            answer = (
                "L’EMSP propose cinq filières : DSER (Digitalisation des Services), "
                "FDIG (Finance Digitale), MDIG (Marketing Digital), LNUM "
                "(Logistique et Numérique) et GARE (Gestion des Activités "
                "Réglementées de l’Économie). Consultez la section « Formations » "
                "de l’accueil pour leurs descriptions."
            )
            if wants_links:
                add_links(
                    {"label": "Voir les formations", "url": "/#formations"},
                    {"label": "Commencer une candidature", "url": "/connexion"},
                )
        elif contains("candidature", "inscription", "inscrire", "postuler"):
            answer = (
                "Pour candidater, cliquez sur « Candidater », connectez-vous ou "
                "créez votre dossier, puis complétez les étapes : informations "
                "personnelles, BAC, tuteurs et documents."
            )
            if wants_links:
                add_links(
                    {"label": "Se connecter", "url": "/connexion"},
                    {"label": "Créer un dossier", "url": "/candidature?inscription=1"},
                )
        elif contains("connexion", "connecter", "mot de passe", "identifiant"):
            answer = (
                "La connexion se fait avec votre adresse email ou votre numéro de "
                "dossier, ainsi que votre mot de passe. En cas de difficulté, "
                "la page Contact permet d’envoyer un message à l’équipe."
            )
            if wants_links:
                add_links(
                    {"label": "Page de connexion", "url": "/connexion"},
                    {"label": "Contacter l’équipe", "url": "/contacts"},
                )
        elif contains("document", "piece", "bulletin", "attestation", "releve", "cni", "cv"):
            answer = (
                "Dans la section « Documents justificatifs », déposez les pièces "
                "demandées : attestation ou relevé du BAC, pièce d’identité, "
                "bulletins, photo et autres documents utiles. Les formats et la "
                "taille autorisés sont indiqués dans le formulaire."
            )
            if wants_links:
                add_links(
                    {"label": "Accéder à la candidature", "url": "/connexion"},
                    {"label": "Créer un dossier", "url": "/candidature?inscription=1"},
                )
        elif contains("convocation", "composition", "concours", "centre", "date"):
            answer = (
                "La rubrique « Ma convocation » présente votre statut, votre numéro "
                "de dossier, la date et le centre de composition lorsqu’ils sont "
                "disponibles. Le bouton de téléchargement permet d’obtenir le PDF."
            )
            if wants_links:
                add_links(
                    {"label": "Accéder à mon dossier", "url": "/connexion"},
                )
        elif contains("contact", "telephone", "email", "adresse", "aide", "joindre"):
            answer = (
                "Vous pouvez joindre l’équipe depuis la page « Contact » en "
                "remplissant le formulaire avec votre nom, votre email, votre objet "
                "et votre message."
            )
            if wants_links:
                add_links(
                    {"label": "Page Contact", "url": "/contacts"},
                    {"label": "Accueil EMSP", "url": "/"},
                )
        elif contains("merci", "super", "parfait"):
            answer = "Avec plaisir ! Je reste disponible pour vous guider."
        else:
            answer = (
                "Je peux vous aider concernant les formations, la candidature, la "
                "connexion, les documents, la convocation ou le contact avec l’EMSP. "
                "Pouvez-vous préciser votre question ?"
            )

        if wants_links and not links:
            add_links(
                {"label": "Accueil EMSP", "url": "/"},
                {"label": "Formations", "url": "/#formations"},
                {"label": "Contact", "url": "/contacts"},
            )

        return jsonify(answer=answer, links=links[:3])

    return app


app = create_app()

if __name__ == '__main__':
    app.run(
        host='0.0.0.0',
        port=int(os.getenv('PORT', '5000')),
        debug=os.getenv('FLASK_DEBUG', 'false').lower() == 'true',
    )
