"""Outils Excel pour l'importation groupée des résultats."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


RESULT_HEADERS = [
    "numero_dossier", "email", "nom", "prenoms",
    "note_francais_compo", "note_math_compo", "note_anglais_compo",
    "note_psycho_compo", "admis_concours", "filiere_formation", "publier",
]


def create_results_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultats"
    sheet.append(RESULT_HEADERS)
    sheet.append([
        "CDT_EXEMPLE", "candidat@exemple.ci", "NOM", "Prénoms",
        12.5, 14, 13, 15, "OUI", "DSER (Digitalisation des Services)", "NON",
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155B3A")
    widths = [22, 30, 18, 24, 22, 20, 20, 24, 18, 42, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    yes_no = DataValidation(type="list", formula1='"OUI,NON"', allow_blank=True)
    sheet.add_data_validation(yes_no)
    yes_no.add("I2:I1000")
    yes_no.add("K2:K1000")

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Structure du fichier de résultats"])
    instructions.append(["Une ligne correspond à un candidat."])
    instructions.append(["numero_dossier est l'identifiant prioritaire pour associer le résultat."])
    instructions.append(["Si le dossier est vide, l'email peut être utilisé comme identifiant."])
    instructions.append(["publier = OUI rend le résultat visible dans l'espace candidat."])
    instructions.append(["Les notes doivent être comprises entre 0 et 20."])
    instructions.column_dimensions["A"].width = 110
    return workbook


def template_bytes():
    output = BytesIO()
    create_results_template().save(output)
    output.seek(0)
    return output


if __name__ == "__main__":
    create_results_template().save("modele_resultats_emsp.xlsx")
